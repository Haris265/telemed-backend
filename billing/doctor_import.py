"""Parse clinic doctor Excel/CSV uploads and create linked doctors."""

from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook

from catalog.models import Clinic, Speciality

from .serializers import create_doctor_for_clinic

User = get_user_model()

IMPORT_HEADERS = [
    "first_name",
    "last_name",
    "email",
    "password",
    "specialities",
    "session_time",
    "consultation_fee",
]


def build_import_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"
    ws.append(IMPORT_HEADERS)
    ws.append(
        [
            "Ali",
            "Khan",
            "ali.khan@example.com",
            "pass12345",
            "General Physician, Cardiology",
            15,
            1000,
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _rows_from_xlsx(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    headers = [_normalize_header(h) for h in header_row]
    out = []
    for raw in rows_iter:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        row = {}
        for i, key in enumerate(headers):
            if not key:
                continue
            row[key] = raw[i] if i < len(raw) else None
        out.append(row)
    return out


def _rows_from_csv(file_obj) -> list[dict]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode("utf-8-sig")
    else:
        text = raw
    reader = csv.DictReader(io.StringIO(text))
    out = []
    for row in reader:
        normalized = {_normalize_header(k): v for k, v in row.items()}
        if all(not str(v or "").strip() for v in normalized.values()):
            continue
        out.append(normalized)
    return out


def parse_doctor_import_file(uploaded) -> list[dict]:
    name = (getattr(uploaded, "name", "") or "").lower()
    if name.endswith(".csv"):
        return _rows_from_csv(uploaded)
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _rows_from_xlsx(uploaded)
    # Sniff: try xlsx then csv
    try:
        uploaded.seek(0)
        return _rows_from_xlsx(uploaded)
    except Exception:
        uploaded.seek(0)
        return _rows_from_csv(uploaded)


def _resolve_speciality_ids(names_raw) -> list[int]:
    if names_raw is None:
        raise ValueError("specialities is required")
    text = str(names_raw).strip()
    if not text:
        raise ValueError("specialities is required")
    names = [n.strip() for n in text.split(",") if n.strip()]
    if not names:
        raise ValueError("specialities is required")
    ids = []
    for name in names:
        spec = Speciality.objects.filter(name__iexact=name, is_active=True).first()
        if not spec:
            raise ValueError(f"Unknown speciality: {name}")
        ids.append(spec.id)
    return list(set(ids))


def _parse_session_time(value) -> int:
    if value is None or str(value).strip() == "":
        return 15
    try:
        n = int(value)
    except (TypeError, ValueError):
        raise ValueError("session_time must be an integer") from None
    if n < 5 or n > 240:
        raise ValueError("session_time must be between 5 and 240")
    return n


def _parse_fee(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError("consultation_fee must be a number") from None


def import_doctors_for_clinic(clinic: Clinic, uploaded) -> dict:
    rows = parse_doctor_import_file(uploaded)
    created = 0
    failed = []
    for idx, row in enumerate(rows, start=2):  # header is row 1
        email = str(row.get("email") or "").strip().lower()
        try:
            first_name = str(row.get("first_name") or "").strip()
            last_name = str(row.get("last_name") or "").strip()
            password = str(row.get("password") or "").strip()
            if not first_name or not last_name:
                raise ValueError("first_name and last_name are required")
            if not email:
                raise ValueError("email is required")
            if len(password) < 6:
                raise ValueError("password must be at least 6 characters")
            if User.objects.filter(email__iexact=email).exists():
                raise ValueError("A user with this email already exists")
            speciality_ids = _resolve_speciality_ids(row.get("specialities"))
            session_time = _parse_session_time(row.get("session_time"))
            fee = _parse_fee(row.get("consultation_fee"))
            create_doctor_for_clinic(
                clinic=clinic,
                first_name=first_name,
                last_name=last_name,
                email=email,
                password=password,
                speciality_ids=speciality_ids,
                session_time=session_time,
                consultation_fee=fee,
            )
            created += 1
        except Exception as exc:
            failed.append({"row": idx, "email": email or None, "error": str(exc)})
    return {"created": created, "failed": failed, "total_rows": len(rows)}
